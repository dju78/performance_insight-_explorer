"""Data ingestion module for Performance Insight Explorer.
Accepts CSV, XLS, XLSX files, detects encoding/sheets, and preserves pristine raw data.
Supports legacy .xls via xlrd and modern .xlsx via openpyxl.
"""
import io
import os
from typing import Tuple, List, Dict, Any, Optional, Union
import pandas as pd
import openpyxl
try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False


class IngestionError(Exception):
    """Custom exception for data ingestion failures."""
    pass


def get_excel_sheet_names(file_source: Union[str, io.BytesIO, bytes, Any]) -> List[str]:
    """Extract sheet names from an Excel file without loading entire data sheets."""
    try:
        if isinstance(file_source, (str, os.PathLike)):
            ext = os.path.splitext(str(file_source))[1].lower()
            if ext == ".xls":
                if HAS_XLRD:
                    wb = xlrd.open_workbook(file_source, on_demand=True)
                    sheet_names = wb.sheet_names()
                    return sheet_names
                else:
                    xl = pd.ExcelFile(file_source)
                    return xl.sheet_names
            else:
                wb = openpyxl.load_workbook(file_source, read_only=True, keep_links=False)
                sheet_names = wb.sheetnames
                wb.close()
                return sheet_names
        elif isinstance(file_source, (io.BytesIO, bytes)):
            bio = file_source if isinstance(file_source, io.BytesIO) else io.BytesIO(file_source)
            try:
                wb = openpyxl.load_workbook(bio, read_only=True, keep_links=False)
                sheet_names = wb.sheetnames
                wb.close()
                return sheet_names
            except Exception:
                bio.seek(0)
                if HAS_XLRD:
                    wb = xlrd.open_workbook(file_contents=bio.read(), on_demand=True)
                    return wb.sheet_names()
                xl = pd.ExcelFile(bio)
                return xl.sheet_names
        else:
            xl = pd.ExcelFile(file_source)
            return xl.sheet_names
    except Exception as e:
        raise IngestionError(f"Failed to read Excel sheet names: {str(e)}") from e


def load_file(
    file_source: Union[str, io.BytesIO, bytes, Any],
    filename: str = "uploaded_file",
    sheet_name: Optional[str] = None
) -> Tuple[pd.DataFrame, List[str], Dict[str, Any]]:
    """Load tabular data from CSV, XLS, or XLSX into a DataFrame and extract metadata.

    Returns:
        (df_raw, list_of_sheets, metadata_dict)
    """
    ext = os.path.splitext(filename)[1].lower()
    sheet_names: List[str] = []
    active_sheet: str = ""

    try:
        if ext in [".xlsx", ".xlsm"]:
            sheet_names = get_excel_sheet_names(file_source)
            if not sheet_names:
                raise IngestionError("Excel file contains no readable worksheets.")

            active_sheet = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]

            if hasattr(file_source, "seek"):
                file_source.seek(0)

            df = pd.read_excel(file_source, sheet_name=active_sheet, engine="openpyxl")

        elif ext == ".xls":
            sheet_names = get_excel_sheet_names(file_source)
            if not sheet_names:
                raise IngestionError("Legacy XLS file contains no readable worksheets.")

            active_sheet = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]

            if hasattr(file_source, "seek"):
                file_source.seek(0)

            engine_to_use = "xlrd" if HAS_XLRD else None
            df = pd.read_excel(file_source, sheet_name=active_sheet, engine=engine_to_use)

        elif ext in [".csv", ".txt"]:
            encodings = ["utf-8", "utf-8-sig", "latin1", "iso-8859-1", "cp1252"]
            df = None
            last_err = None

            for enc in encodings:
                try:
                    if hasattr(file_source, "seek"):
                        file_source.seek(0)
                    df = pd.read_csv(file_source, encoding=enc, sep=None, engine="python")
                    break
                except Exception as e:
                    last_err = e
                    continue

            if df is None:
                raise IngestionError(f"Unable to parse CSV with standard encodings. Error: {last_err}")

            sheet_names = ["CSV_Default"]
            active_sheet = "CSV_Default"

        else:
            raise IngestionError(f"Unsupported file extension '{ext}'. Supported formats: .csv, .xlsx, .xls")

        if df is None:
            raise IngestionError("Failed to load dataset: DataFrame is None.")

        file_size_bytes = 0
        if isinstance(file_source, (str, os.PathLike)) and os.path.exists(file_source):
            file_size_bytes = os.path.getsize(file_source)
        elif hasattr(file_source, "getbuffer"):
            file_size_bytes = len(file_source.getbuffer())
        elif hasattr(file_source, "getvalue"):
            file_size_bytes = len(file_source.getvalue())

        metadata = {
            "filename": filename,
            "file_format": ext.replace(".", "").upper(),
            "file_size_bytes": file_size_bytes,
            "sheet_names": sheet_names,
            "active_sheet": active_sheet,
            "row_count": int(len(df)),
            "column_count": int(len(df.columns)),
            "columns": list(df.columns)
        }

        return df.copy(deep=True), sheet_names, metadata

    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Error reading '{filename}': {str(e)}") from e


from src.profiling import profile_dataset


def generate_dataset_profile(df: pd.DataFrame) -> Dict[str, Any]:
    """Generate profile dictionary compatible with UI and test assertions."""
    prof = profile_dataset(df)
    cols_dict = {}
    for col_info in prof.get("columns", []):
        cols_dict[col_info["name"]] = col_info
    prof["columns"] = cols_dict
    return prof


def ingest_file(
    file_source: Union[str, io.BytesIO, bytes, Any],
    filename: Optional[str] = None,
    sheet_name: Optional[str] = None
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, Any]]:
    """Convenience pipeline returning (raw_df, clean_df, profile).

    When an Excel workbook is supplied, ``sheet_name`` is forwarded to
    ``load_file`` so the analyst-selected worksheet is loaded instead of
    silently defaulting to the first worksheet.
    """
    fname = filename
    if fname is None:
        if isinstance(file_source, (str, os.PathLike)):
            fname = os.path.basename(str(file_source))
        elif hasattr(file_source, "name"):
            fname = getattr(file_source, "name", "uploaded_file.csv")
        else:
            fname = "uploaded_file.csv"

    raw_df, sheets, meta = load_file(file_source, fname, sheet_name=sheet_name)
    clean_df = raw_df.copy(deep=True)
    profile = generate_dataset_profile(clean_df)
    # Preserve file/sheet metadata in the profile for downstream UI/export use.
    profile["filename"] = meta.get("filename", fname)
    profile["active_sheet"] = meta.get("active_sheet", "")
    profile["sheet_names"] = meta.get("sheet_names", sheets)
    profile["row_count"] = meta.get("row_count", len(clean_df))
    profile["col_count"] = meta.get("column_count", len(clean_df.columns))
    return raw_df, clean_df, profile
