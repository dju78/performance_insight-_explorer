"""Comprehensive unit & structural verification tests for 16:9 executive PowerPoint exporter with Presenter Notes."""
import os
import pytest
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.powerpoint import generate_powerpoint_presentation, generate_interview_powerpoint
from src.export import generate_powerpoint_deck


def test_powerpoint_widescreen_16_9_dimensions(tmp_path):
    """Verify generated presentation is strictly 16:9 widescreen (13.333 x 7.5 inches)."""
    out_file = str(tmp_path / "deck_16x9.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15, "row_granularity": "Periodic snapshot"}
    qa = {"critical_count": 0, "warning_count": 1, "health_score": 85.0, "issues": []}
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary={},
        trend_summary=None,
        comparison_summary=None,
        insights=[],
        recommendations={},
        limitations=[],
        assumptions=[],
        assessment_context={"assessment_question": "Test Question", "target_audience": "Senior Leadership"},
        include_appendix=False
    )
    
    prs = Presentation(path)
    assert round(prs.slide_width.inches, 2) == 13.33
    assert round(prs.slide_height.inches, 2) == 7.50
    assert len(prs.slides) == 6


def test_powerpoint_chart_images_on_slide_4(tmp_path):
    """Verify Slide 4 contains embedded picture shapes for time trends and cohort comparisons."""
    out_file = str(tmp_path / "deck_charts.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15}
    qa = {"critical_count": 0, "warning_count": 0, "health_score": 100.0, "issues": []}
    
    df = pd.DataFrame({
        "Period": ["2025-01", "2025-02", "2025-03", "2025-04"],
        "Team": ["East", "West", "North", "South"],
        "Volume": [100, 150, 120, 180]
    })
    mappings = {"Period": "period", "Team": "group", "Volume": "completed"}
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary={},
        trend_summary=None,
        comparison_summary=None,
        insights=[],
        recommendations={},
        limitations=[],
        assumptions=[],
        clean_df=df,
        confirmed_mappings=mappings
    )
    
    prs = Presentation(path)
    slide4 = prs.slides[3] # Slide 4 (0-indexed 3)
    
    pictures = [s for s in slide4.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) >= 2, f"Expected at least 2 embedded chart pictures on Slide 4, found {len(pictures)}"


def test_powerpoint_kpi_cards_and_diagnostic_slide_3(tmp_path):
    """Verify Slide 3 contains live KPI numbers and does not show 'No calculated KPIs available'."""
    out_file = str(tmp_path / "deck_kpi.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15}
    qa = {"critical_count": 0, "warning_count": 0, "health_score": 100.0, "issues": []}
    
    kpis = {
        "summary_kpis": {
            "target_achievement_pct": {"name": "Target Achievement", "value": 92.7, "unit": "%"},
            "productivity": {"name": "Productivity", "value": 14.72, "unit": "cases/FTE-period"},
            "utilisation_pct": {"name": "Utilisation", "value": 85.8, "unit": "%"},
            "backlog_change": {"name": "Backlog Change", "value": 785.0, "unit": "cases"}
        }
    }
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary=kpis,
        trend_summary=None,
        comparison_summary=None,
        insights=[],
        recommendations={},
        limitations=[],
        assumptions=[]
    )
    
    prs = Presentation(path)
    slide3 = prs.slides[2] # Slide 3
    slide3_text = " ".join([s.text_frame.text for s in slide3.shapes if s.has_text_frame])
    
    assert "No calculated KPIs available" not in slide3_text
    assert "92.7%" in slide3_text
    assert "14.72" in slide3_text
    assert "85.8%" in slide3_text
    assert "+785" in slide3_text or "785" in slide3_text
    assert "Executive Performance Diagnostic" in slide3_text


def test_powerpoint_approved_only_insights_slide_5(tmp_path):
    """Verify Slide 5 strictly includes approved insights and avoids unreviewed/pending ones."""
    out_file = str(tmp_path / "deck_insights.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15}
    qa = {"critical_count": 0, "warning_count": 0, "health_score": 100.0, "issues": []}
    
    insights = [
        {"id": "INS-1", "category": "Target", "title": "Approved Insight Title", "finding": "Confirmed target gap.", "status": "approved"},
        {"id": "INS-2", "category": "Capacity", "title": "Pending Unverified Finding", "finding": "Unverified assertion.", "status": "pending"}
    ]
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary={},
        trend_summary=None,
        comparison_summary=None,
        insights=insights,
        recommendations={},
        limitations=[],
        assumptions=[]
    )
    
    prs = Presentation(path)
    slide5 = prs.slides[4] # Slide 5
    slide5_text = " ".join([s.text_frame.text for s in slide5.shapes if s.has_text_frame])
    
    assert "Approved Insight Title" in slide5_text
    assert "Pending Unverified Finding" not in slide5_text


def test_powerpoint_action_table_and_limitations_slide_6(tmp_path):
    """Verify Slide 6 contains the structured action matrix table and analytical limitations."""
    out_file = str(tmp_path / "deck_actions.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15}
    qa = {"critical_count": 0, "warning_count": 0, "health_score": 100.0, "issues": []}
    
    recs = {
        "Capacity": [
            {"id": "REC-1", "title": "Rebalance FTE", "action": "Reallocate 2 FTE from Central to West", "owner": "Ops Manager", "timeframe": "Week 1", "expected_impact": "Reduce SLA breaches", "status": "approved"}
        ]
    }
    limits = [{"limitation": "Snapshot Granularity", "impact": "Aggregated data"}]
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary={},
        trend_summary=None,
        comparison_summary=None,
        insights=[],
        recommendations=recs,
        limitations=limits,
        assumptions=[]
    )
    
    prs = Presentation(path)
    slide6 = prs.slides[5] # Slide 6
    
    tables = [s.table for s in slide6.shapes if s.has_table]
    assert len(tables) == 1, "Expected 1 Action Matrix table on Slide 6"
    
    tbl = tables[0]
    cell_texts = [tbl.cell(r, c).text_frame.text for r in range(len(tbl.rows)) for c in range(len(tbl.columns))]
    full_table_text = " ".join(cell_texts)
    
    assert "ACTION ITEM & INTERVENTION" in full_table_text
    assert "Rebalance FTE" in full_table_text
    assert "Ops Manager" in full_table_text
    assert "Week 1" in full_table_text


def test_powerpoint_appendix_extension_10_slides(tmp_path):
    """Verify include_appendix=True creates exactly 10 slides."""
    out_file = str(tmp_path / "deck_10_slides.pptx")
    meta = {"filename": "test.csv", "row_count": 60, "column_count": 15}
    qa = {"critical_count": 0, "warning_count": 0, "health_score": 100.0, "issues": []}
    
    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary={},
        trend_summary=None,
        comparison_summary=None,
        insights=[],
        recommendations={},
        limitations=[],
        assumptions=[],
        include_appendix=True
    )
    
    prs = Presentation(path)
    assert len(prs.slides) == 10


def test_powerpoint_presenter_notes_on_all_slides(tmp_path):
    """Verify all slides have real, non-empty, first-person presenter notes with Q&A defence cues."""
    out_file = str(tmp_path / "deck_with_notes.pptx")
    meta = {
        "filename": "performance_insight_explorer_test_data.csv",
        "row_count": 60,
        "column_count": 15,
        "row_granularity": "Periodic operational snapshot"
    }
    qa = {
        "critical_count": 0,
        "warning_count": 1,
        "health_score": 85.0,
        "issues": [
            {"issue_type": "missing_values", "description": "1 missing value in Output_Target"},
            {"issue_type": "case_inconsistency", "description": "North Operations and north operations casing variation"},
            {"issue_type": "outliers", "description": "Processing duration outlier 70.6 days"},
            {"issue_type": "zero_denominator", "description": "Available_FTE = 0 in August"}
        ]
    }
    kpis = {
        "summary_kpis": {
            "target_achievement_pct": {"name": "Target Achievement", "value": 92.7, "unit": "%"},
            "target_variance": {"name": "Target Variance", "value": -738.0, "unit": "units"},
            "productivity": {"name": "Productivity", "value": 14.72, "unit": "cases/FTE-period"},
            "utilisation_pct": {"name": "Utilisation", "value": 85.8, "unit": "%"},
            "backlog_change": {"name": "Backlog Change", "value": 785.0, "unit": "cases"}
        }
    }
    insights = [
        {"id": "ins-1", "title": "Backlog Accumulation in West", "finding": "Demand outpaced capacity leading to 210 case increase.", "status": "approved", "pillar": "Queue Flow"},
        {"id": "ins-2", "title": "Rejected Noise", "finding": "Unverified assertion.", "status": "rejected", "pillar": "Other"}
    ]
    recs = {
        "Capacity": [
            {"id": "rec-1", "title": "FTE Rebalancing", "action": "Deploy floating FTE to West Operations", "owner": "Ops Lead", "timeframe": "Immediate (1-2 Weeks)", "expected_impact": "Reduce backlog by 150 cases", "status": "approved"},
            {"id": "rec-2", "title": "Rejected Rec", "action": "Do nothing", "owner": "Nobody", "timeframe": "Never", "status": "rejected"}
        ]
    }
    ctx = {
        "assessment_question": "Evaluate multi-team operational throughput and diagnose queue bottlenecks.",
        "target_audience": "Chief Operating Officer & Senior Leadership"
    }

    df = pd.DataFrame({
        "Period": ["2025-01", "2025-02", "2025-03", "2025-04", "2025-05", "2025-06", "2025-07", "2025-08", "2025-09", "2025-10", "2025-11", "2025-12"],
        "Service_Team": ["East", "West", "Central", "South", "North", "north operations", "East", "West", "Central", "South", "North", "West"],
        "Demand": [846, 882, 915, 948, 890, 860, 830, 805, 780, 752, 790, 810],
        "Completed": [800, 820, 850, 900, 820, 800, 780, 750, 730, 710, 740, 760]
    })
    mappings = {"Period": "period", "Service_Team": "team", "Demand": "demand", "Completed": "completed"}

    path = generate_powerpoint_presentation(
        output_filepath=out_file,
        project_metadata=meta,
        qa_report=qa,
        kpi_summary=kpis,
        trend_summary=None,
        comparison_summary=None,
        insights=insights,
        recommendations=recs,
        limitations=[{"limitation": "Periodic Snapshot", "impact": "Aggregated data"}],
        assumptions=[{"assumption": "Standard calendar", "area": "Capacity"}],
        assessment_context=ctx,
        clean_df=df,
        confirmed_mappings=mappings,
        include_appendix=True
    )

    prs = Presentation(path)
    assert len(prs.slides) == 10

    placeholders = ["TODO", "TBC", "N/A", "Insert figure", "Generic finding", "Lorem ipsum"]

    total_word_count = 0
    for idx, slide in enumerate(prs.slides):
        assert slide.notes_slide is not None, f"Slide {idx+1} is missing notes_slide"
        tf = slide.notes_slide.notes_text_frame
        assert tf is not None, f"Slide {idx+1} notes_text_frame is None"
        notes_text = tf.text.strip()
        assert len(notes_text) > 50, f"Slide {idx+1} presenter notes are too short or empty: '{notes_text}'"

        # Check for first person voice
        assert any(term in notes_text for term in ["I ", "I'", "My ", "my "]), f"Slide {idx+1} notes lack first-person candidate voice: {notes_text[:100]}"

        # Check for panel Q&A section
        assert "POSSIBLE FOLLOW-UP QUESTIONS" in notes_text, f"Slide {idx+1} is missing panel Q&A section"
        assert "Q:" in notes_text and "A:" in notes_text, f"Slide {idx+1} Q&A section missing Q/A prompts"

        # Check for prohibited placeholders
        for ph in placeholders:
            assert ph not in notes_text, f"Slide {idx+1} contains forbidden placeholder '{ph}'"

        words = len(notes_text.split())
        total_word_count += words

    # Verify slide 1 notes contain live metadata
    s1_notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "60 observations" in s1_notes
    assert "15 variables" in s1_notes
    assert "Periodic operational snapshot" in s1_notes or "periodic operational snapshot" in s1_notes

    # Verify slide 2 notes contain health score & anomalies
    s2_notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "85.0" in s2_notes or "85" in s2_notes
    assert "North Operations" in s2_notes

    # Verify slide 3 notes contain live KPIs
    s3_notes = prs.slides[2].notes_slide.notes_text_frame.text
    assert "92.7%" in s3_notes
    assert "738" in s3_notes
    assert "14.72" in s3_notes
    assert "85.8%" in s3_notes
    assert "785" in s3_notes

    # Verify slide 4 notes contain trend wave & casing caution
    s4_notes = prs.slides[3].notes_slide.notes_text_frame.text
    assert "948" in s4_notes or "peak" in s4_notes.lower()
    assert "752" in s4_notes or "trough" in s4_notes.lower()

    # Verify rejected items are absent from slide 5 & 6 notes
    s5_notes = prs.slides[4].notes_slide.notes_text_frame.text
    assert "Backlog Accumulation in West" in s5_notes
    assert "Rejected Noise" not in s5_notes

    s6_notes = prs.slides[5].notes_slide.notes_text_frame.text
    assert "FTE Rebalancing" in s6_notes
    assert "Rejected Rec" not in s6_notes

    assert total_word_count >= 1000, f"Total presenter notes word count across 10 slides should be >= 1000, got {total_word_count}"
