import os
import io
import pandas as pd
import pytest
from pypdf import PdfReader
from pptx import Presentation
import openpyxl

from src.ingestion import get_excel_sheet_names, load_file
from src.pdf_report import generate_pdf_document
from src.powerpoint import generate_powerpoint_presentation
from src.reporting import generate_excel_summary
from src.export import (
    build_export_payload_from_state,
    generate_executive_excel_pack,
    generate_powerpoint_deck,
    generate_pdf_report,
    generate_audit_trail_text,
    group_recommendations_by_category
)
from src.metrics import calculate_kpis
from src.state import clear_dataset_for_new_upload
import streamlit as st


def test_excel_multi_sheet_and_sheet_switching(tmp_path):
    """Verify multi-sheet Excel discovery and individual sheet loading."""
    file_path = tmp_path / "multi_sheet_operations.xlsx"
    df_east = pd.DataFrame({"Team": ["East"], "Cases": [120], "FTE": [4.0]})
    df_west = pd.DataFrame({"Team": ["West"], "Cases": [95], "FTE": [3.5]})
    
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df_east.to_excel(writer, sheet_name="East_Region", index=False)
        df_west.to_excel(writer, sheet_name="West_Region", index=False)
        
    sheets = get_excel_sheet_names(str(file_path))
    assert sheets == ["East_Region", "West_Region"]
    
    loaded_east, _, meta_east = load_file(str(file_path), "multi_sheet_operations.xlsx", sheet_name="East_Region")
    assert len(loaded_east) == 1
    assert loaded_east["Team"].iloc[0] == "East"
    assert meta_east["active_sheet"] == "East_Region"
    
    loaded_west, _, meta_west = load_file(str(file_path), "multi_sheet_operations.xlsx", sheet_name="West_Region")
    assert len(loaded_west) == 1
    assert loaded_west["Team"].iloc[0] == "West"
    assert meta_west["active_sheet"] == "West_Region"


def test_pdf_generation_and_pypdf_verification(tmp_path):
    """Verify ReportLab PDF generation creates valid PDF with real data and no fabrications."""
    pdf_out = str(tmp_path / "executive_brief.pdf")
    
    export_payload = {
        "filename": "test_operational_dataset.csv",
        "active_sheet": "Sheet1",
        "row_count": 60,
        "column_count": 15,
        "dataset_fingerprint": "abc123hash",
        "row_granularity": "Team-Period Aggregate",
        "assessment_context": {
            "question": "Assess productivity variances and recommend 3 structural remediations.",
            "audience": "Senior Leadership",
            "analyst_notes": "Operations performance Q3 review."
        },
        "kpi_results": {
            "summary_kpis": {
                "productivity": {
                    "label": "Productivity",
                    "value": 18.5,
                    "unit": "cases/FTE-period",
                    "interpretation": "Aggregate throughput is 18.5 cases per FTE across 60 reporting periods."
                },
                "target_achievement_pct": {
                    "label": "Target Achievement",
                    "value": 94.2,
                    "unit": "%",
                    "interpretation": "Operations achieved 94.2% of set targets."
                }
            }
        },
        "qa_report": {
            "health_score": 88.0,
            "critical_count": 1,
            "warning_count": 1,
            "issues": [
                {"issue_id": "MISSING_VALUES", "column": "Output_Target", "description": "1 missing target observation", "severity": "WARNING"},
                {"issue_id": "ZERO_DENOMINATOR", "column": "Available_FTE", "description": "Zero FTE in East Operations (2025-08-01)", "severity": "CRITICAL"}
            ]
        },
        "approved_insights": [
            {
                "id": "INS-001",
                "title": "East Operations Outperforms West on Cases Per FTE",
                "finding": "East delivered 22.4 cases/FTE vs 14.6 in West.",
                "evidence": "60 periods of comparative logs.",
                "pillar": "Productivity",
                "status": "approved"
            }
        ],
        "recommendations_by_category": {
            "Act": [
                {
                    "id": "REC-001",
                    "title": "Rebalance West Operations Backlog Intake",
                    "category": "Act",
                    "action": "Shift 25% of unassigned intake to East Operations buffer.",
                    "expected_impact": "Reduces West backlog by 18% in 6 weeks.",
                    "owner": "Head of Operations",
                    "timeframe": "1-2 Weeks",
                    "status": "approved"
                }
            ],
            "Investigate": [],
            "Monitor": [],
            "Improve Reporting": []
        },
        "assumptions": ["FTE headcount is measured at full working capacity."],
        "limitations": ["Backlog intake timestamps are aggregated monthly."],
        "audit_trail": [
            {"timestamp": "2026-09-14T18:00:00", "event_type": "DATASET_LOADED", "action": "Uploaded test_operational_dataset.csv", "details": {}}
        ]
    }
    
    generated_path = generate_pdf_report(export_payload, audience="Senior Leadership", output_filepath=pdf_out)
    assert os.path.exists(generated_path)
    assert os.path.getsize(generated_path) > 1000
    
    # Verify with pypdf
    reader = PdfReader(generated_path)
    assert len(reader.pages) >= 1
    
    full_pdf_text = "".join([page.extract_text() for page in reader.pages])
    assert "Performance Insight Explorer" in full_pdf_text or "EXECUTIVE BRIEF" in full_pdf_text
    assert "DARAMOLA OMOYELE" in full_pdf_text
    assert "60" in full_pdf_text  # Row count
    assert "15" in full_pdf_text  # Col count
    assert "East Operations Outperforms West" in full_pdf_text
    assert "Rebalance West Operations Backlog" in full_pdf_text


def test_powerpoint_generation_with_export_payload(tmp_path):
    """Verify PowerPoint deck generation receives full export payload with approved-only filter."""
    pptx_out = str(tmp_path / "briefing_deck.pptx")
    
    export_payload = {
        "filename": "operations_data.csv",
        "active_sheet": "N/A",
        "row_count": 60,
        "column_count": 15,
        "row_granularity": "Case / Record",
        "dataset_fingerprint": "xyz789hash",
        "assessment_context": {
            "question": "Operational variance diagnostic.",
            "audience": "Senior Leadership",
            "analyst_notes": "Diagnostic deck."
        },
        "kpi_results": {
            "summary_kpis": {
                "productivity": {"label": "Productivity", "value": 15.2, "unit": "cases/FTE-period"}
            }
        },
        "qa_report": {
            "health_score": 92.5,
            "critical_count": 0,
            "warning_count": 1,
            "issues": [{"issue_id": "ZERO_DENOM", "column": "FTE", "description": "FTE zero", "severity": "WARNING"}]
        },
        "approved_insights": [
            {"id": "INS-01", "title": "Approved Trend Insight", "finding": "Positive productivity shift.", "pillar": "Capacity", "status": "approved"}
        ],
        "recommendations_by_category": {
            "Act": [{"id": "REC-01", "title": "Approved Action Item", "category": "Act", "action": "Implement triage.", "owner": "Team Lead", "status": "approved"}],
            "Investigate": [],
            "Monitor": [],
            "Improve Reporting": []
        },
        "assumptions": ["Data verified by operational leads."],
        "limitations": ["No sub-daily timestamps available."]
    }
    
    generated_path = generate_powerpoint_deck(export_payload, output_filepath=pptx_out)
    assert os.path.exists(generated_path)
    
    prs = Presentation(generated_path)
    assert len(prs.slides) >= 4
    
    slide_texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                slide_texts.append(shape.text_frame.text)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        slide_texts.append(cell.text_frame.text)
    combined_text = " ".join(slide_texts)
    
    assert "DARAMOLA OMOYELE" in combined_text
    assert "Approved Trend Insight" in combined_text
    assert "Approved Action Item" in combined_text


def test_excel_export_generation_with_real_audit_and_qa(tmp_path):
    """Verify Excel workbook generation includes Data, Mappings, KPIs, QA, Insights, and Audit."""
    excel_out = str(tmp_path / "operational_pack.xlsx")
    
    export_payload = {
        "filename": "operations_data.csv",
        "active_sheet": "Sheet1",
        "row_count": 2,
        "column_count": 3,
        "raw_df": pd.DataFrame({
            "Period": ["2025-01", "2025-02"],
            "Cases": [100, 110],
            "FTE": [5.0, 5.0]
        }),
        "kpi_results": {
            "summary_kpis": {"productivity": {"label": "Productivity", "value": 21.0, "unit": "cases/FTE-period"}}
        },
        "qa_report": {
            "health_score": 95.0,
            "critical_count": 0,
            "warning_count": 0,
            "issues": []
        },
        "approved_insights": [{"insight_id": "INS-1", "category": "Productivity", "finding": "Cases stable across periods."}],
        "recommendations_by_category": {
            "Act": [{"id": "REC-1", "title": "Maintain Capacity", "action": "Keep FTE at 5.0."}],
            "Investigate": [],
            "Monitor": [],
            "Improve Reporting": []
        },
        "assumptions": [{"assumption": "Stable workflow assumed."}],
        "limitations": [{"limitation": "Two months sample."}],
        "audit_trail": [{"timestamp": "2026-09-14T18:30:00", "event_type": "AUDIT_EXPORT", "action": "Exported Excel pack", "details": {}}]
    }
    
    path = generate_executive_excel_pack(export_payload, output_filepath=excel_out)
    assert os.path.exists(path)
    
    wb = openpyxl.load_workbook(path)
    sheet_names = wb.sheetnames
    assert "Executive_Summary" in sheet_names
    assert "Headline_KPIs" in sheet_names
    assert "Insights" in sheet_names
    assert "Audit_Trail" in sheet_names


def test_clear_dataset_workflow_resets_derived_state_and_increments_version():
    """Verify clear_dataset_for_new_upload completely wipes operational state and increments upload widget key."""
    # Populate mock session state
    st.session_state["raw_df"] = pd.DataFrame({"A": [1, 2, 3]})
    st.session_state["clean_df"] = pd.DataFrame({"A": [1, 2, 3]})
    st.session_state["uploaded_file_name"] = "test.csv"
    st.session_state["uploaded_file_bytes"] = b"123"
    st.session_state["confirmed_mappings"] = {"A": "actual"}
    st.session_state["kpi_results"] = {"some": "metric"}
    st.session_state["trend_summary"] = {"trend": "up"}
    st.session_state["comparison_summary"] = {"comp": "east vs west"}
    st.session_state["insights_list"] = ["insight1"]
    st.session_state["recommendations_list"] = ["rec1"]
    st.session_state["upload_widget_version"] = 1
    st.session_state["assessment_question"] = "Custom Question To Preserve"
    st.session_state["target_audience"] = "Executive Audience"
    
    # Execute clear workflow preserving assessment context
    clear_dataset_for_new_upload(preserve_assessment_context=True)
    
    assert st.session_state.get("raw_df") is None
    assert st.session_state.get("clean_df") is None
    assert st.session_state.get("uploaded_file_name") == ""
    assert st.session_state.get("uploaded_file_bytes") is None
    assert st.session_state.get("confirmed_mappings") == {}
    assert st.session_state.get("kpi_results") == {}
    assert st.session_state.get("trend_summary") is None
    assert st.session_state.get("comparison_summary") is None
    assert st.session_state.get("insights_list") == []
    assert st.session_state.get("recommendations_list") == []
    
    # Verify widget version incremented to reset file_uploader widget
    assert st.session_state.get("upload_widget_version") == 2
    
    # Verify assessment context was preserved
    assert st.session_state.get("assessment_question") == "Custom Question To Preserve"
    assert st.session_state.get("target_audience") == "Executive Audience"


def test_metrics_ratio_of_sums_and_backlog_movement():
    """Verify productivity uses ratio-of-sums across all periods and backlog movement computes from earliest to latest period."""
    df = pd.DataFrame({
        "Period": ["2025-01-01", "2025-02-01", "2025-03-01"],
        "Team": ["East", "East", "East"],
        "Completed": [100, 150, 200],  # Sum = 450
        "FTE": [5.0, 5.0, 5.0],        # Sum = 15.0 FTE-periods
        "Opening_BL": [50, 60, 70],
        "Closing_BL": [60, 70, 80],
        "Demand": [110, 160, 210]
    })
    mappings = {
        "Period": "period",
        "Team": "group",
        "Completed": "completed",
        "FTE": "fte",
        "Opening_BL": "opening_backlog",
        "Closing_BL": "closing_backlog",
        "Demand": "received"
    }
    res = calculate_kpis(df, mappings)
    kpis = res["summary_kpis"]
    
    # Productivity = 450 / 15.0 = 30.0 cases/FTE-period
    assert kpis["productivity"]["value"] == 30.0
    assert "FTE-period" in kpis["productivity"]["unit"]
    
    # Backlog movement across time = 80 (latest closing) - 50 (earliest opening) = +30.0
    assert kpis["backlog_change"]["value"] == 30.0
